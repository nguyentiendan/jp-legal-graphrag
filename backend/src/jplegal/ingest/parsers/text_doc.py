"""Parser for non-e-Gov-XML uploads: PDF / DOCX / TXT / MD.

Split articles by lines starting with 「第◯条」; the caption is the 「（…）」 line right above.
If no 条 is found → split by paragraph (~max_chars) into pseudo-articles c1, c2, ...
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from jplegal.core.models import Article, LawMeta, Paragraph, ParsedLaw
from jplegal.ingest.normalize import NUM_CLASS, article_key_from_title

_ART_LINE_RE = re.compile(rf"^\s*(第{NUM_CLASS}条(?:の{NUM_CLASS})*)[\s　]*(.*)$")
_CAPTION_RE = re.compile(r"^\s*[（(][^）)]{1,40}[）)]\s*$")
_PARA_RE = re.compile(r"^\s*([２-９2-9]|[１-９1-9][０-９0-9])[\s　]+(.*)$")


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    data = path.read_bytes()
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return "\n".join(page.get_text("text") for page in doc)
    if suffix == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs)
    for enc in ("utf-8", "utf-8-sig", "cp932", "euc-jp"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_text_document(text: str, law_id: str, title: str, max_chars: int = 1500) -> ParsedLaw:
    lines = [ln.rstrip() for ln in text.replace("\r", "").split("\n")]
    meta = LawMeta(law_id=law_id, title=title, law_type="Upload", source="upload")
    articles: list[Article] = []
    current: Article | None = None
    pending_caption = ""
    path: list[str] = []
    seen_keys: dict[str, int] = {}

    def flush_para(buf: list[str], label: str, num: int) -> None:
        if current is not None and buf:
            current.paragraphs.append(Paragraph(num=num, label=label, text="".join(s.strip() for s in buf)))

    buf: list[str] = []
    label, pnum = "", 1
    for ln in lines:
        if not ln.strip():
            continue
        if re.match(rf"^\s*第{NUM_CLASS}[章節款]", ln) and len(ln) < 60:
            flush_para(buf, label, pnum)
            buf = []
            path = [ln.strip()]
            continue
        if _CAPTION_RE.match(ln):
            pending_caption = ln.strip()
            continue
        m = _ART_LINE_RE.match(ln)
        if m:
            flush_para(buf, label, pnum)
            key = article_key_from_title(m.group(1)) or f"x{len(articles)}"
            if key in seen_keys:  # duplicate article number (e.g. 附則) → append suffix
                seen_keys[key] += 1
                key = f"{key}-{seen_keys[key]}"
            else:
                seen_keys[key] = 0
            current = Article(law_id=law_id, article_key=key, title=m.group(1), caption=pending_caption,
                              path=list(path), order=len(articles))
            articles.append(current)
            pending_caption = ""
            buf, label, pnum = [m.group(2)] if m.group(2) else [], "", 1
            continue
        pm = _PARA_RE.match(ln)
        if pm and current is not None:
            flush_para(buf, label, pnum)
            pnum += 1
            buf, label = [pm.group(2)], pm.group(1)
            continue
        if current is not None:  # leading text before 第一条 is skipped
            buf.append(ln)
    flush_para(buf, label, pnum)

    if not articles:  # no 条 structure → split by paragraph
        body = "\n".join(lines)
        paras = [p.strip() for p in re.split(r"\n\s*\n|(?<=。)\n", body) if p.strip()]
        group: list[str] = []
        size = 0
        for p in paras + [None]:  # type: ignore[list-item]
            if p is None or (size and size + len(p) > max_chars):
                if group:
                    i = len(articles) + 1
                    articles.append(Article(law_id=law_id, article_key=f"c{i}", title=f"{title} ({i})",
                                            paragraphs=[Paragraph(num=1, label="", text="\n".join(group))],
                                            order=i - 1))
                group, size = [], 0
            if p is not None:
                group.append(p)
                size += len(p)
    return ParsedLaw(meta=meta, articles=articles)
